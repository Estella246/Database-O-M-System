from __future__ import annotations

import json
from datetime import date, datetime
from io import BytesIO
from typing import Any

import psycopg
from psycopg.errors import UndefinedTable

from fastapi import APIRouter, File, Form, HTTPException, UploadFile

from config import (
    DUTY_ROTATION_ROSTER_KINDS,
    _DUTY_EXTRAS_SCHEMA_HINT,
    _HOLIDAY_SCHEMA_HINT,
)
from database import db_conn
from models import (
    DutyCalendarPutPayload,
    DutyCalendarSlotPayload,
    DutyRotationPutPayload,
    DutySiteOnCallPutPayload,
    DutyRlOnCallPutPayload,
    HolidayConfigPutPayload,
)

_DUTY_CALENDAR_KINDS = ("kernel", "control", "public_cloud", "poc", "research_version")


def _normalize_calendar_kind(raw: str) -> str:
    kind = str(raw or "").strip()
    if kind not in _DUTY_CALENDAR_KINDS:
        raise HTTPException(
            status_code=400,
            detail="kind 须为 kernel、control、public_cloud、poc 或 research_version",
        )
    return kind


def _parse_calendar_date_key(raw: str) -> str:
    text = str(raw or "").strip()
    try:
        return date.fromisoformat(text).isoformat()
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="date 须为 YYYY-MM-DD") from exc
from utils import duty_month_bounds as _duty_month_bounds
from leave_duty_effect import sync_leave_duty_status
from utils.logging_config import audit_log
from whitelist_policy import duty_roster_edit_rl_only, whitelist_field_levels, whitelist_permission_level

router = APIRouter(prefix="/api/duty", tags=["duty"])


def _require_duty_roster_edit(conn: psycopg.Connection, operator_id: str, *, rl_only: bool = False) -> None:
    wl = whitelist_field_levels(conn, operator_id.strip() or "")
    level = whitelist_permission_level(wl, "duty_roster_edit")
    if level == "hidden":
        raise HTTPException(status_code=403, detail="无编辑权限")
    if not rl_only and duty_roster_edit_rl_only(wl):
        raise HTTPException(status_code=403, detail="无编辑权限")


def _normalize_duty_shift(raw) -> str | None:
    val = str(raw or "").strip().lower()
    if val in ("", "full", "全天"):
        return "full"
    if val in ("night", "晚班"):
        return "night"
    return None


def _normalize_excel_date_string(raw: str) -> str | None:
    text = str(raw or "").strip()
    if not text:
        return None
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:10]
    for sep in ("/", "-", "."):
        parts = [p.strip() for p in text.split(sep) if p.strip()]
        if len(parts) < 3:
            continue
        try:
            year = int(parts[0])
            month = int(parts[1])
            day = int(parts[2])
            return date(year, month, day).isoformat()
        except (TypeError, ValueError):
            continue
    return None


def _parse_excel_date_key(
    raw,
    row_idx: int,
    errors: list[dict],
    *,
    month_prefix: str | None = None,
) -> str | None:
    if raw is None or str(raw).strip() == "":
        errors.append({"row": row_idx, "field": "日期", "message": "必填字段不能为空"})
        return None
    dk: str | None = None
    if isinstance(raw, datetime):
        dk = raw.date().isoformat()
    elif isinstance(raw, date):
        dk = raw.isoformat()
    elif isinstance(raw, (int, float)) and not isinstance(raw, bool):
        from openpyxl.utils.datetime import from_excel

        try:
            converted = from_excel(raw)
        except (TypeError, ValueError, OverflowError):
            converted = None
        if isinstance(converted, datetime):
            dk = converted.date().isoformat()
        elif isinstance(converted, date):
            dk = converted.isoformat()
    else:
        dk = _normalize_excel_date_string(str(raw).strip())
    if not dk or len(dk) != 10:
        errors.append({"row": row_idx, "field": "日期", "message": "日期格式无效，须为 YYYY-MM-DD"})
        return None
    if month_prefix is not None and not dk.startswith(month_prefix):
        errors.append({"row": row_idx, "field": "日期", "message": f"日期须属于当月（{month_prefix}）"})
        return None
    return dk


def _parse_duty_calendar_excel(
    file_content: bytes,
    *,
    year: int,
    month: int,
) -> tuple[dict[str, list[dict[str, str]]], list[dict]]:
    from openpyxl import load_workbook

    month_prefix = f"{year}-{month:02d}-"
    wb = load_workbook(BytesIO(file_content))
    ws = wb.active

    headers: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        header_val = ws.cell(row=1, column=col).value
        if header_val:
            headers[str(header_val).strip()] = col

    expected_headers = ["日期", "账号", "姓名", "班次"]
    errors: list[dict] = []
    for h in expected_headers:
        if h not in headers:
            errors.append({"row": 1, "field": "表头", "message": f"缺少必填列：{h}"})
    if errors:
        return {}, errors

    days: dict[str, list[dict[str, str]]] = {}
    for row_idx in range(2, ws.max_row + 1):
        account_raw = ws.cell(row=row_idx, column=headers["账号"]).value
        account = str(account_raw or "").strip()
        if not account:
            continue

        date_raw = ws.cell(row=row_idx, column=headers["日期"]).value
        dk = _parse_excel_date_key(date_raw, row_idx, errors, month_prefix=month_prefix)
        if dk is None:
            continue

        shift_raw = ws.cell(row=row_idx, column=headers["班次"]).value
        shift = _normalize_duty_shift(shift_raw)
        if shift is None:
            errors.append({"row": row_idx, "field": "班次", "message": "班次须为 全天 或 晚班"})
            continue

        name_raw = ws.cell(row=row_idx, column=headers["姓名"]).value
        user_name = str(name_raw or "").strip()

        days.setdefault(dk, []).append(
            {
                "account": account,
                "user_name": user_name,
                "shift": shift,
                "_row_idx": row_idx,
            }
        )

    return days, errors


def _validate_duty_calendar_days_payload(
    *,
    year: int,
    month: int,
    days: dict[str, list[dict[str, Any]]],
) -> None:
    prefix = f"{year}-{month:02d}-"
    for dk, slots in days.items():
        if not isinstance(dk, str) or not dk.startswith(prefix):
            raise HTTPException(status_code=400, detail=f"日期键须属于当月: {dk}")
        if not isinstance(slots, list):
            raise HTTPException(status_code=400, detail="班次列表格式无效")
        for slot in slots:
            if not isinstance(slot, dict):
                raise HTTPException(status_code=400, detail="班次项格式无效")
            sh = str(slot.get("shift") or "full")
            if sh not in ("full", "night"):
                raise HTTPException(status_code=400, detail="shift 须为 full 或 night")
            acc = str(slot.get("account") or "").strip()
            if not acc:
                raise HTTPException(status_code=400, detail="account 不能为空")


def _apply_duty_calendar_days(
    conn: psycopg.Connection,
    *,
    kind: str,
    year: int,
    month: int,
    days: dict[str, list[dict[str, Any]]],
    operator_id: str,
) -> dict[str, int]:
    """仅处理 payload 中的日期：按 (account, shift) 差量增删，保留未改行的 last_accept_at。"""
    _validate_duty_calendar_days_payload(year=year, month=month, days=days)
    deleted = inserted = updated = kept = 0
    for dk, slots in days.items():
        desired: list[dict[str, str]] = []
        for slot in slots:
            if not isinstance(slot, dict):
                continue
            desired.append(
                {
                    "account": str(slot.get("account") or "").strip(),
                    "user_name": str(slot.get("user_name") or "").strip(),
                    "shift": str(slot.get("shift") or "full"),
                }
            )
        existing = conn.execute(
            """
            SELECT id, account, user_name, shift
            FROM duty_calendar_assignment
            WHERE table_kind = %s AND duty_date = %s::date
            ORDER BY id
            """,
            (kind, dk),
        ).fetchall()
        unmatched = [dict(row) for row in existing]
        to_insert: list[dict[str, str]] = []
        for slot in desired:
            match_idx = next(
                (
                    i
                    for i, row in enumerate(unmatched)
                    if str(row.get("account") or "") == slot["account"]
                    and str(row.get("shift") or "") == slot["shift"]
                ),
                None,
            )
            if match_idx is None:
                to_insert.append(slot)
                continue
            row = unmatched.pop(match_idx)
            if str(row.get("user_name") or "") != slot["user_name"]:
                conn.execute(
                    """
                    UPDATE duty_calendar_assignment
                    SET user_name = %s, updated_by = %s, updated_at = NOW()
                    WHERE id = %s
                    """,
                    (slot["user_name"], operator_id, int(row["id"])),
                )
                updated += 1
            else:
                kept += 1
        for row in unmatched:
            conn.execute("DELETE FROM duty_calendar_assignment WHERE id = %s", (int(row["id"]),))
            deleted += 1
        for slot in to_insert:
            conn.execute(
                """
                INSERT INTO duty_calendar_assignment (
                  table_kind, duty_date, account, user_name, shift, updated_by, updated_at
                )
                VALUES (%s, %s::date, %s, %s, %s, %s, NOW())
                """,
                (kind, dk, slot["account"], slot["user_name"], slot["shift"], operator_id),
            )
            inserted += 1
    return {"deleted": deleted, "inserted": inserted, "updated": updated, "kept": kept}


def _replace_duty_calendar_month(
    conn: psycopg.Connection,
    *,
    kind: str,
    year: int,
    month: int,
    days: dict[str, list[dict[str, Any]]],
    operator_id: str,
) -> None:
    """整月覆盖（导入用）：先删当月再插入。"""
    start, end = _duty_month_bounds(year, month)
    _validate_duty_calendar_days_payload(year=year, month=month, days=days)

    conn.execute(
        """
        DELETE FROM duty_calendar_assignment
        WHERE table_kind = %s AND duty_date >= %s AND duty_date < %s
        """,
        (kind, start, end),
    )
    for dk, slots in days.items():
        for slot in slots:
            if not isinstance(slot, dict):
                continue
            conn.execute(
                """
                INSERT INTO duty_calendar_assignment (
                  table_kind, duty_date, account, user_name, shift, updated_by, updated_at
                )
                VALUES (%s, %s::date, %s, %s, %s, %s, NOW())
                """,
                (
                    kind,
                    dk,
                    str(slot.get("account") or "").strip(),
                    str(slot.get("user_name") or "").strip(),
                    str(slot.get("shift") or "full"),
                    operator_id,
                ),
            )


def _normalize_day_type(v) -> str:
    raw = str(v or "").strip()
    if raw in ("workday", "工作日"):
        return "workday"
    if raw in ("weekend_holiday", "周末节假日"):
        return "weekend_holiday"
    raise HTTPException(status_code=400, detail="day_type 仅允许 workday(工作日) 或 weekend_holiday(周末节假日)")


def _normalize_duty_status(raw) -> str:
    st = str(raw or "active").strip()
    if st in ("active", "当值"):
        return "active"
    if st in ("inactive", "非当值"):
        return "inactive"
    raise HTTPException(status_code=400, detail="status 须为 active/inactive 或 当值/非当值")


def _validate_duty_rotation_put_lists(lists: dict) -> None:
    unknown = set(lists.keys()) - set(DUTY_ROTATION_ROSTER_KINDS)
    if unknown:
        raise HTTPException(status_code=400, detail=f"未知 roster_kind: {sorted(unknown)}")
    for kind in DUTY_ROTATION_ROSTER_KINDS:
        items = lists.get(kind)
        if items is None:
            continue
        if not isinstance(items, list):
            raise HTTPException(status_code=400, detail=f"{kind} 须为数组")
        for slot in items:
            if not isinstance(slot, dict):
                raise HTTPException(status_code=400, detail="轮值项格式无效")
            acc = str(slot.get("account") or "").strip()
            if not acc:
                raise HTTPException(status_code=400, detail=f"{kind} 中存在空的 account")
            st = str(slot.get("status") or "active").strip()
            if st not in ("active", "inactive", "当值", "非当值"):
                raise HTTPException(status_code=400, detail="status 须为 active/inactive 或 当值/非当值")


@router.get("/calendar")
def get_duty_calendar(year: int, month: int, operator_id: str = "demo_001") -> dict:
    _ = operator_id
    start, end = _duty_month_bounds(year, month)
    out_kernel: dict[str, list[dict[str, str]]] = {}
    out_control: dict[str, list[dict[str, str]]] = {}
    out_public_cloud: dict[str, list[dict[str, str]]] = {}
    out_poc: dict[str, list[dict[str, str]]] = {}
    out_research_version: dict[str, list[dict[str, str]]] = {}
    try:
        with db_conn() as conn:
            rows = conn.execute(
                """
                SELECT table_kind, duty_date, account, user_name, shift
                FROM duty_calendar_assignment
                WHERE duty_date >= %s AND duty_date < %s
                ORDER BY table_kind, duty_date, id
                """,
                (start, end),
            ).fetchall()
    except UndefinedTable as exc:
        raise HTTPException(
            status_code=503,
            detail="值班日历表未创建，请在数据库执行 db/migrations/0016_duty_calendar_assignment.sql",
        ) from exc
    for row in rows:
        raw_d = row["duty_date"]
        if isinstance(raw_d, date):
            dk = raw_d.isoformat()
        else:
            dk = str(raw_d)[:10]
        item = {
            "account": str(row["account"] or ""),
            "user_name": str(row["user_name"] or ""),
            "shift": str(row["shift"] or "full"),
        }
        kind = str(row["table_kind"] or "")
        if kind == "kernel":
            out_kernel.setdefault(dk, []).append(item)
        elif kind == "control":
            out_control.setdefault(dk, []).append(item)
        elif kind == "public_cloud":
            out_public_cloud.setdefault(dk, []).append(item)
        elif kind == "poc":
            out_poc.setdefault(dk, []).append(item)
        elif kind == "research_version":
            out_research_version.setdefault(dk, []).append(item)
    return {
        "year": year,
        "month": month,
        "kernel": out_kernel,
        "control": out_control,
        "public_cloud": out_public_cloud,
        "poc": out_poc,
        "research_version": out_research_version,
    }


@router.put("/calendar")
def put_duty_calendar(payload: DutyCalendarPutPayload) -> dict:
    kind = _normalize_calendar_kind(payload.kind)
    op = payload.operator_id.strip() or "admin"
    try:
        with db_conn() as conn:
            _require_duty_roster_edit(conn, op)
            diff = _apply_duty_calendar_days(
                conn,
                kind=kind,
                year=payload.year,
                month=payload.month,
                days=payload.days,
                operator_id=op,
            )
            conn.commit()
    except UndefinedTable as exc:
        raise HTTPException(
            status_code=503,
            detail="值班日历表未创建，请在数据库执行 db/migrations/0016_duty_calendar_assignment.sql",
        ) from exc
    audit_log(
        "duty.calendar.update",
        operator=op,
        kind=kind,
        year=payload.year,
        month=payload.month,
        day_count=len(payload.days),
        inserted=diff["inserted"],
        deleted=diff["deleted"],
        updated=diff["updated"],
        kept=diff["kept"],
    )
    return {
        "ok": True,
        "kind": kind,
        "year": payload.year,
        "month": payload.month,
        "diff": diff,
    }


@router.post("/calendar/slot")
def add_duty_calendar_slot(payload: DutyCalendarSlotPayload) -> dict:
    """单条新增排班；不影响同日其他人，不覆盖 last_accept_at。"""
    kind = _normalize_calendar_kind(payload.kind)
    op = payload.operator_id.strip() or "admin"
    duty_date = _parse_calendar_date_key(payload.date)
    shift = _normalize_duty_shift(payload.shift)
    if shift is None:
        raise HTTPException(status_code=400, detail="shift 须为 full 或 night")
    account = str(payload.account or "").strip()
    if not account:
        raise HTTPException(status_code=400, detail="account 不能为空")
    user_name = str(payload.user_name or "").strip()
    try:
        with db_conn() as conn:
            _require_duty_roster_edit(conn, op)
            row = conn.execute(
                """
                INSERT INTO duty_calendar_assignment (
                  table_kind, duty_date, account, user_name, shift, updated_by, updated_at
                )
                VALUES (%s, %s::date, %s, %s, %s, %s, NOW())
                RETURNING id
                """,
                (kind, duty_date, account, user_name, shift, op),
            ).fetchone()
            conn.commit()
    except UndefinedTable as exc:
        raise HTTPException(
            status_code=503,
            detail="值班日历表未创建，请在数据库执行 db/migrations/0016_duty_calendar_assignment.sql",
        ) from exc
    slot_id = int(row["id"]) if row else 0
    audit_log(
        "duty.calendar.slot.add",
        operator=op,
        kind=kind,
        date=duty_date,
        account=account,
        shift=shift,
        id=slot_id,
    )
    return {
        "ok": True,
        "id": slot_id,
        "kind": kind,
        "date": duty_date,
        "account": account,
        "user_name": user_name,
        "shift": shift,
    }


@router.delete("/calendar/slot")
def delete_duty_calendar_slot(payload: DutyCalendarSlotPayload) -> dict:
    """单条删除排班；仅删匹配的一条，不影响同日其他人。"""
    kind = _normalize_calendar_kind(payload.kind)
    op = payload.operator_id.strip() or "admin"
    duty_date = _parse_calendar_date_key(payload.date)
    shift = _normalize_duty_shift(payload.shift)
    if shift is None:
        raise HTTPException(status_code=400, detail="shift 须为 full 或 night")
    account = str(payload.account or "").strip()
    if not account:
        raise HTTPException(status_code=400, detail="account 不能为空")
    try:
        with db_conn() as conn:
            _require_duty_roster_edit(conn, op)
            row = conn.execute(
                """
                SELECT id
                FROM duty_calendar_assignment
                WHERE table_kind = %s AND duty_date = %s::date AND account = %s AND shift = %s
                ORDER BY id
                LIMIT 1
                """,
                (kind, duty_date, account, shift),
            ).fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="未找到对应排班")
            slot_id = int(row["id"])
            conn.execute("DELETE FROM duty_calendar_assignment WHERE id = %s", (slot_id,))
            conn.commit()
    except HTTPException:
        raise
    except UndefinedTable as exc:
        raise HTTPException(
            status_code=503,
            detail="值班日历表未创建，请在数据库执行 db/migrations/0016_duty_calendar_assignment.sql",
        ) from exc
    audit_log(
        "duty.calendar.slot.delete",
        operator=op,
        kind=kind,
        date=duty_date,
        account=account,
        shift=shift,
        id=slot_id,
    )
    return {
        "ok": True,
        "id": slot_id,
        "kind": kind,
        "date": duty_date,
        "account": account,
        "shift": shift,
    }


@router.post("/calendar/import")
async def import_duty_calendar(
    file: UploadFile = File(...),
    operator_id: str = Form(...),
    kind: str = Form(...),
    year: int = Form(...),
    month: int = Form(...),
) -> dict:
    """批量导入月历值班表（整月覆盖）。"""
    op = operator_id.strip() or "demo_001"
    kind = kind.strip()
    if kind not in ("kernel", "control", "public_cloud", "poc", "research_version"):
        raise HTTPException(status_code=400, detail="kind 须为 kernel、control、public_cloud、poc 或 research_version")
    if year < 2000 or year > 2100 or month < 1 or month > 12:
        raise HTTPException(status_code=400, detail="year/month 无效")

    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 格式文件")

    try:
        content = await file.read()
        days, parse_errors = _parse_duty_calendar_excel(content, year=year, month=month)
        if parse_errors:
            raise HTTPException(
                status_code=400,
                detail=json.dumps({"success": False, "error_type": "validation_failed", "errors": parse_errors}),
            )
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="文件无法解析，请检查文件格式") from exc

    total_slots = sum(len(v) for v in days.values())

    try:
        with db_conn() as conn:
            _require_duty_roster_edit(conn, op)

            accounts_in_file: set[str] = set()
            account_rows: list[tuple[int, str]] = []
            for slots in days.values():
                for slot in slots:
                    if not isinstance(slot, dict):
                        continue
                    acc = str(slot.get("account") or "").strip()
                    if acc:
                        accounts_in_file.add(acc)
                        account_rows.append((int(slot.get("_row_idx") or 0), acc))

            if accounts_in_file:
                rows = conn.execute(
                    "SELECT account, user_name FROM user_account WHERE account = ANY(%s)",
                    (list(accounts_in_file),),
                ).fetchall()
                known = {str(r["account"] or "").strip(): str(r["user_name"] or "").strip() for r in rows}
                validation_errors: list[dict] = []
                for row_idx, acc in account_rows:
                    if acc not in known:
                        validation_errors.append(
                            {"row": row_idx, "field": "账号", "message": f"账号不存在：{acc}"}
                        )
                if validation_errors:
                    raise HTTPException(
                        status_code=400,
                        detail=json.dumps(
                            {"success": False, "error_type": "validation_failed", "errors": validation_errors}
                        ),
                    )

                for slots in days.values():
                    for slot in slots:
                        if not isinstance(slot, dict):
                            continue
                        acc = str(slot.get("account") or "").strip()
                        if acc and not str(slot.get("user_name") or "").strip():
                            slot["user_name"] = known.get(acc, "")
                        slot.pop("_row_idx", None)
            else:
                for slots in days.values():
                    for slot in slots:
                        if isinstance(slot, dict):
                            slot.pop("_row_idx", None)

            _replace_duty_calendar_month(
                conn,
                kind=kind,
                year=year,
                month=month,
                days=days,
                operator_id=op,
            )
            conn.commit()
    except HTTPException:
        raise
    except UndefinedTable as exc:
        raise HTTPException(
            status_code=503,
            detail="值班日历表未创建，请在数据库执行 db/migrations/0016_duty_calendar_assignment.sql",
        ) from exc

    return {
        "success": True,
        "kind": kind,
        "year": year,
        "month": month,
        "total": total_slots,
        "message": f"导入成功，共 {total_slots} 条排班",
    }


@router.get("/holidays")
def get_holiday_config(year: int, month: int, operator_id: str = "demo_001") -> dict:
    _ = operator_id
    start, end = _duty_month_bounds(year, month)
    out: dict[str, str] = {}
    try:
        with db_conn() as conn:
            rows = conn.execute(
                """
                SELECT holiday_date, day_type
                FROM holiday_day_config
                WHERE holiday_date >= %s AND holiday_date < %s
                ORDER BY holiday_date
                """,
                (start, end),
            ).fetchall()
    except UndefinedTable as exc:
        raise HTTPException(status_code=503, detail=f"节假日配置表未就绪：{_HOLIDAY_SCHEMA_HINT}") from exc
    for row in rows:
        raw_d = row.get("holiday_date")
        dk = raw_d.isoformat() if isinstance(raw_d, date) else str(raw_d)[:10]
        out[dk] = _normalize_day_type(row.get("day_type"))
    return {"year": year, "month": month, "days": out}


@router.put("/holidays")
def put_holiday_config(payload: HolidayConfigPutPayload) -> dict:
    op = payload.operator_id.strip() or "admin"
    start, end = _duty_month_bounds(payload.year, payload.month)
    prefix = f"{payload.year}-{payload.month:02d}-"
    normalized_days: dict[str, str] = {}
    for dk, day_type in payload.days.items():
        if not isinstance(dk, str) or not dk.startswith(prefix):
            raise HTTPException(status_code=400, detail=f"日期键须属于当月: {dk}")
        normalized_days[dk] = _normalize_day_type(day_type)
    try:
        with db_conn() as conn:
            _require_duty_roster_edit(conn, op)
            conn.execute(
                """
                DELETE FROM holiday_day_config
                WHERE holiday_date >= %s AND holiday_date < %s
                """,
                (start, end),
            )
            for dk, day_type in normalized_days.items():
                conn.execute(
                    """
                    INSERT INTO holiday_day_config (holiday_date, day_type, updated_by, updated_at)
                    VALUES (%s::date, %s, %s, NOW())
                    """,
                    (dk, day_type, op),
                )
            conn.commit()
    except UndefinedTable as exc:
        raise HTTPException(status_code=503, detail=f"节假日配置表未就绪：{_HOLIDAY_SCHEMA_HINT}") from exc
    audit_log(
        "duty.holidays.update",
        operator=op,
        year=payload.year,
        month=payload.month,
        day_count=len(normalized_days),
    )
    return {"ok": True, "year": payload.year, "month": payload.month}


@router.get("/rotation")
def get_duty_rotation(operator_id: str = "demo_001") -> dict:
    _ = operator_id
    out: dict[str, list[dict[str, str]]] = {k: [] for k in DUTY_ROTATION_ROSTER_KINDS}
    try:
        with db_conn() as conn:
            sync_leave_duty_status(conn)
            conn.commit()
            rows = conn.execute(
                """
                SELECT roster_kind, position, account, user_name, status, last_accept_at
                FROM duty_rotation_entry
                ORDER BY roster_kind, position
                """
            ).fetchall()
    except UndefinedTable as exc:
        raise HTTPException(status_code=503, detail=f"轮值表未就绪：{_DUTY_EXTRAS_SCHEMA_HINT}") from exc
    for row in rows:
        kind = str(row["roster_kind"] or "")
        if kind not in out:
            continue
        out[kind].append(
            {
                "account": str(row["account"] or ""),
                "user_name": str(row["user_name"] or ""),
                "status": str(row["status"] or "active"),
                "last_accept_at": str(row["last_accept_at"] or ""),
            }
        )
    return out


def _rotation_entry_preserve_key(roster_kind: str, account: str) -> tuple[str, str]:
    return (str(roster_kind or "").strip(), str(account or "").strip())


def _load_rotation_dispatch_preserve_map(conn: psycopg.Connection) -> dict[tuple[str, str], dict[str, Any]]:
    rows = conn.execute(
        """
        SELECT roster_kind, account, last_accept_at,
               last_dispatch_at, last_dispatch_ticket_no,
               last_dispatch_node_key, last_dispatch_rule
        FROM duty_rotation_entry
        """
    ).fetchall()
    out: dict[tuple[str, str], dict[str, Any]] = {}
    for row in rows:
        key = _rotation_entry_preserve_key(str(row.get("roster_kind") or ""), str(row.get("account") or ""))
        if not key[1]:
            continue
        out[key] = dict(row)
    return out


@router.put("/rotation")
def put_duty_rotation(payload: DutyRotationPutPayload) -> dict:
    op = payload.operator_id.strip() or "admin"
    lists = payload.lists if isinstance(payload.lists, dict) else {}
    _validate_duty_rotation_put_lists(lists)
    try:
        with db_conn() as conn:
            _require_duty_roster_edit(conn, op)
            preserve_map = _load_rotation_dispatch_preserve_map(conn)
            conn.execute("DELETE FROM duty_rotation_entry")
            for kind in DUTY_ROTATION_ROSTER_KINDS:
                items = lists.get(kind) or []
                for pos, slot in enumerate(items):
                    if not isinstance(slot, dict):
                        continue
                    acct = str(slot.get("account") or "").strip()
                    preserved = preserve_map.get(_rotation_entry_preserve_key(kind, acct))
                    if preserved:
                        last_accept_at = str(preserved.get("last_accept_at") or "").strip()[:64]
                        last_dispatch_at = preserved.get("last_dispatch_at")
                        last_dispatch_ticket_no = str(preserved.get("last_dispatch_ticket_no") or "").strip()[:32]
                        last_dispatch_node_key = str(preserved.get("last_dispatch_node_key") or "").strip()[:64]
                        last_dispatch_rule = preserved.get("last_dispatch_rule")
                        if not isinstance(last_dispatch_rule, dict):
                            last_dispatch_rule = {}
                    else:
                        last_accept_at = ""
                        last_dispatch_at = None
                        last_dispatch_ticket_no = ""
                        last_dispatch_node_key = ""
                        last_dispatch_rule = {}
                    conn.execute(
                        """
                        INSERT INTO duty_rotation_entry (
                          roster_kind, position, account, user_name, status, last_accept_at,
                          last_dispatch_at, last_dispatch_ticket_no, last_dispatch_node_key,
                          last_dispatch_rule, updated_by, updated_at
                        )
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb, %s, NOW())
                        """,
                        (
                            kind,
                            pos,
                            acct,
                            str(slot.get("user_name") or "").strip(),
                            "active",
                            last_accept_at,
                            last_dispatch_at,
                            last_dispatch_ticket_no,
                            last_dispatch_node_key,
                            psycopg.types.json.Jsonb(last_dispatch_rule),
                            op,
                        ),
                    )
            sync_leave_duty_status(conn, updated_by=op)
            conn.commit()
    except UndefinedTable as exc:
        raise HTTPException(status_code=503, detail=f"轮值表未就绪：{_DUTY_EXTRAS_SCHEMA_HINT}") from exc
    audit_log("duty.rotation.update", operator=op, roster_kinds=len(DUTY_ROTATION_ROSTER_KINDS))
    return {"ok": True}


@router.get("/site-oncall")
def get_duty_site_oncall(operator_id: str = "demo_001") -> dict:
    _ = operator_id
    rows_out: list[dict[str, str]] = []
    try:
        with db_conn() as conn:
            sync_leave_duty_status(conn)
            conn.commit()
            rows = conn.execute(
                """
                SELECT position, site_name, account, user_name, status, last_accept_at
                FROM duty_site_oncall_row
                ORDER BY position
                """
            ).fetchall()
    except UndefinedTable as exc:
        raise HTTPException(status_code=503, detail=f"局点值班表未就绪：{_DUTY_EXTRAS_SCHEMA_HINT}") from exc
    for row in rows:
        rows_out.append(
            {
                "site_name": str(row["site_name"] or ""),
                "account": str(row["account"] or ""),
                "user_name": str(row["user_name"] or ""),
                "status": str(row["status"] or "active"),
                "last_accept_at": str(row["last_accept_at"] or ""),
            }
        )
    return {"rows": rows_out}


@router.put("/site-oncall")
def put_duty_site_oncall(payload: DutySiteOnCallPutPayload) -> dict:
    op = payload.operator_id.strip() or "admin"
    for i, row in enumerate(payload.rows):
        if not isinstance(row, dict):
            raise HTTPException(status_code=400, detail="行格式无效")
        site = str(row.get("site_name") or "").strip()
        acc = str(row.get("account") or "").strip()
        if not site or not acc:
            raise HTTPException(status_code=400, detail=f"第 {i + 1} 行须含 site_name 与 account")
        st = str(row.get("status") or "active").strip()
        _normalize_duty_status(st)
    try:
        with db_conn() as conn:
            _require_duty_roster_edit(conn, op)
            conn.execute("DELETE FROM duty_site_oncall_row")
            for pos, row in enumerate(payload.rows):
                if not isinstance(row, dict):
                    continue
                conn.execute(
                    """
                    INSERT INTO duty_site_oncall_row (
                      position, site_name, account, user_name, status, last_accept_at, updated_by, updated_at
                    )
                    VALUES (%s, %s, %s, %s, %s, %s, %s, NOW())
                    """,
                    (
                        pos,
                        str(row.get("site_name") or "").strip(),
                        str(row.get("account") or "").strip(),
                        str(row.get("user_name") or "").strip(),
                        "active",
                        str(row.get("last_accept_at") or "").strip()[:64],
                        op,
                    ),
                )
            sync_leave_duty_status(conn, updated_by=op)
            conn.commit()
    except UndefinedTable as exc:
        raise HTTPException(status_code=503, detail=f"局点值班表未就绪：{_DUTY_EXTRAS_SCHEMA_HINT}") from exc
    audit_log("duty.site_oncall.update", operator=op, count=len(payload.rows))
    return {"ok": True, "count": len(payload.rows)}


@router.get("/rl-oncall")
def get_duty_rl_oncall(operator_id: str = "demo_001") -> dict:
    _ = operator_id
    rows_out: list[dict] = []
    try:
        with db_conn() as conn:
            rows = conn.execute(
                """
                SELECT duty_date, primary_account, primary_user_name, primary_phone,
                       backup_account, backup_user_name, backup_phone
                FROM duty_rl_oncall_row
                ORDER BY duty_date DESC
                """
            ).fetchall()
    except UndefinedTable as exc:
        raise HTTPException(status_code=503, detail=f"RL 值班表未就绪：{_DUTY_EXTRAS_SCHEMA_HINT}") from exc
    for row in rows:
        raw_d = row["duty_date"]
        dk = raw_d.isoformat() if isinstance(raw_d, date) else str(raw_d)[:10]
        rows_out.append(
            {
                "duty_date": dk,
                "primary": {
                    "account": str(row["primary_account"] or ""),
                    "user_name": str(row["primary_user_name"] or ""),
                    "phone": str(row["primary_phone"] or ""),
                },
                "backup": {
                    "account": str(row["backup_account"] or ""),
                    "user_name": str(row["backup_user_name"] or ""),
                    "phone": str(row["backup_phone"] or ""),
                },
            }
        )
    return {"rows": rows_out}


_RL_ONCALL_IMPORT_HEADERS = (
    "日期",
    "主值班账号",
    "主值班姓名",
    "主值班手机",
    "备值班账号",
    "备值班姓名",
    "备值班手机",
)


def _excel_cell_text(raw) -> str:
    if raw is None:
        return ""
    if isinstance(raw, bool):
        return str(raw).strip()
    if isinstance(raw, int):
        return str(raw)
    if isinstance(raw, float):
        if raw == int(raw):
            return str(int(raw))
        return str(raw).strip()
    return str(raw).strip()


def _parse_rl_oncall_excel(file_content: bytes) -> tuple[list[dict[str, Any]], list[dict]]:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(file_content))
    ws = wb.active

    headers: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        header_val = ws.cell(row=1, column=col).value
        if header_val:
            headers[str(header_val).strip()] = col

    errors: list[dict] = []
    for h in _RL_ONCALL_IMPORT_HEADERS:
        if h not in headers:
            errors.append({"row": 1, "field": "表头", "message": f"缺少必填列：{h}"})
    if errors:
        return [], errors

    rows_out: list[dict[str, Any]] = []
    seen_dates: set[str] = set()
    for row_idx in range(2, ws.max_row + 1):
        primary_account = _excel_cell_text(ws.cell(row=row_idx, column=headers["主值班账号"]).value)
        if not primary_account:
            continue

        date_raw = ws.cell(row=row_idx, column=headers["日期"]).value
        dk = _parse_excel_date_key(date_raw, row_idx, errors)
        if dk is None:
            continue
        if dk in seen_dates:
            errors.append({"row": row_idx, "field": "日期", "message": f"重复日期：{dk}"})
            continue
        seen_dates.add(dk)

        primary_name = _excel_cell_text(ws.cell(row=row_idx, column=headers["主值班姓名"]).value)
        primary_phone = _excel_cell_text(ws.cell(row=row_idx, column=headers["主值班手机"]).value)
        backup_account = _excel_cell_text(ws.cell(row=row_idx, column=headers["备值班账号"]).value)
        backup_name = _excel_cell_text(ws.cell(row=row_idx, column=headers["备值班姓名"]).value)
        backup_phone = _excel_cell_text(ws.cell(row=row_idx, column=headers["备值班手机"]).value)

        rows_out.append(
            {
                "duty_date": dk,
                "primary": {
                    "account": primary_account,
                    "user_name": primary_name,
                    "phone": primary_phone,
                },
                "backup": {
                    "account": backup_account,
                    "user_name": backup_name,
                    "phone": backup_phone,
                },
                "_row_idx": row_idx,
            }
        )

    return rows_out, errors


@router.post("/rl-oncall/import")
async def import_duty_rl_oncall(
    file: UploadFile = File(...),
    operator_id: str = Form(...),
) -> dict:
    """批量导入 RL 值班表：按日期覆盖（文件中出现的日期覆盖库中同日记录，其它日期保留）。"""
    op = operator_id.strip() or "demo_001"

    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 格式文件")

    try:
        content = await file.read()
        rows, parse_errors = _parse_rl_oncall_excel(content)
        if parse_errors:
            raise HTTPException(
                status_code=400,
                detail=json.dumps({"success": False, "error_type": "validation_failed", "errors": parse_errors}),
            )
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="文件无法解析，请检查文件格式") from exc

    try:
        with db_conn() as conn:
            _require_duty_roster_edit(conn, op, rl_only=True)

            accounts_in_file: set[str] = set()
            for row in rows:
                pa = str((row.get("primary") or {}).get("account") or "").strip()
                ba = str((row.get("backup") or {}).get("account") or "").strip()
                if pa:
                    accounts_in_file.add(pa)
                if ba:
                    accounts_in_file.add(ba)

            known: dict[str, dict[str, str]] = {}
            if accounts_in_file:
                db_rows = conn.execute(
                    """
                    SELECT account, user_name, contact_phone
                    FROM user_account
                    WHERE account = ANY(%s)
                    """,
                    (list(accounts_in_file),),
                ).fetchall()
                known = {
                    str(r["account"] or "").strip(): {
                        "user_name": str(r["user_name"] or "").strip(),
                        "phone": str(r["contact_phone"] or "").strip(),
                    }
                    for r in db_rows
                }

            validation_errors: list[dict] = []
            for row in rows:
                row_idx = int(row.get("_row_idx") or 0)
                pri = row.get("primary") if isinstance(row.get("primary"), dict) else {}
                bak = row.get("backup") if isinstance(row.get("backup"), dict) else {}
                pa = str(pri.get("account") or "").strip()
                ba = str(bak.get("account") or "").strip()
                if pa and pa not in known:
                    validation_errors.append(
                        {"row": row_idx, "field": "主值班账号", "message": f"账号不存在：{pa}"}
                    )
                if ba and ba not in known:
                    validation_errors.append(
                        {"row": row_idx, "field": "备值班账号", "message": f"账号不存在：{ba}"}
                    )
                if pa in known:
                    if not str(pri.get("user_name") or "").strip():
                        pri["user_name"] = known[pa]["user_name"]
                    if not str(pri.get("phone") or "").strip():
                        pri["phone"] = known[pa]["phone"]
                if ba and ba in known:
                    if not str(bak.get("user_name") or "").strip():
                        bak["user_name"] = known[ba]["user_name"]
                    if not str(bak.get("phone") or "").strip():
                        bak["phone"] = known[ba]["phone"]
                if pa and not str(pri.get("phone") or "").strip():
                    validation_errors.append(
                        {"row": row_idx, "field": "主值班手机", "message": "主值班手机不能为空"}
                    )
                if ba and not str(bak.get("phone") or "").strip():
                    validation_errors.append(
                        {"row": row_idx, "field": "备值班手机", "message": "已填备值班账号时手机不能为空"}
                    )
                row["primary"] = pri
                row["backup"] = bak

            if validation_errors:
                raise HTTPException(
                    status_code=400,
                    detail=json.dumps(
                        {"success": False, "error_type": "validation_failed", "errors": validation_errors}
                    ),
                )

            for row in rows:
                dk = str(row.get("duty_date") or "").strip()[:10]
                pri = row.get("primary") if isinstance(row.get("primary"), dict) else {}
                bak = row.get("backup") if isinstance(row.get("backup"), dict) else {}
                conn.execute("DELETE FROM duty_rl_oncall_row WHERE duty_date = %s::date", (dk,))
                conn.execute(
                    """
                    INSERT INTO duty_rl_oncall_row (
                      duty_date,
                      primary_account, primary_user_name, primary_phone,
                      backup_account, backup_user_name, backup_phone,
                      updated_by, updated_at
                    )
                    VALUES (%s::date, %s, %s, %s, %s, %s, %s, %s, NOW())
                    """,
                    (
                        dk,
                        str(pri.get("account") or "").strip(),
                        str(pri.get("user_name") or "").strip(),
                        str(pri.get("phone") or "").strip()[:32],
                        str(bak.get("account") or "").strip(),
                        str(bak.get("user_name") or "").strip(),
                        str(bak.get("phone") or "").strip()[:32],
                        op,
                    ),
                )
            conn.commit()
    except HTTPException:
        raise
    except UndefinedTable as exc:
        raise HTTPException(status_code=503, detail=f"RL 值班表未就绪：{_DUTY_EXTRAS_SCHEMA_HINT}") from exc

    total = len(rows)
    audit_log("duty.rl_oncall.import", operator=op, count=total)
    return {
        "success": True,
        "total": total,
        "message": f"导入成功，共 {total} 条排班（按日期覆盖）",
    }


@router.put("/rl-oncall")
def put_duty_rl_oncall(payload: DutyRlOnCallPutPayload) -> dict:
    op = payload.operator_id.strip() or "admin"
    seen_dates: set[str] = set()
    for i, row in enumerate(payload.rows):
        if not isinstance(row, dict):
            raise HTTPException(status_code=400, detail="行格式无效")
        dk = str(row.get("duty_date") or "").strip()[:10]
        if not dk or len(dk) != 10:
            raise HTTPException(status_code=400, detail=f"第 {i + 1} 行 duty_date 无效")
        if dk in seen_dates:
            raise HTTPException(status_code=400, detail=f"重复日期: {dk}")
        seen_dates.add(dk)
        pri = row.get("primary") if isinstance(row.get("primary"), dict) else {}
        bak = row.get("backup") if isinstance(row.get("backup"), dict) else {}
        pa = str(pri.get("account") or "").strip()
        pp = str(pri.get("phone") or "").strip()
        if not pa or not pp:
            raise HTTPException(status_code=400, detail=f"日期 {dk} 的主值班须填写 account 与 phone")
        ba = str(bak.get("account") or "").strip()
        bp = str(bak.get("phone") or "").strip()
        if ba and not bp:
            raise HTTPException(status_code=400, detail=f"日期 {dk} 的备值班已选人须填写 phone")
    try:
        with db_conn() as conn:
            _require_duty_roster_edit(conn, op, rl_only=True)
            conn.execute("DELETE FROM duty_rl_oncall_row")
            for row in payload.rows:
                if not isinstance(row, dict):
                    continue
                dk = str(row.get("duty_date") or "").strip()[:10]
                pri = row.get("primary") if isinstance(row.get("primary"), dict) else {}
                bak = row.get("backup") if isinstance(row.get("backup"), dict) else {}
                conn.execute(
                    """
                    INSERT INTO duty_rl_oncall_row (
                      duty_date,
                      primary_account, primary_user_name, primary_phone,
                      backup_account, backup_user_name, backup_phone,
                      updated_by, updated_at
                    )
                    VALUES (%s::date, %s, %s, %s, %s, %s, %s, %s, NOW())
                    """,
                    (
                        dk,
                        str(pri.get("account") or "").strip(),
                        str(pri.get("user_name") or "").strip(),
                        str(pri.get("phone") or "").strip()[:32],
                        str(bak.get("account") or "").strip(),
                        str(bak.get("user_name") or "").strip(),
                        str(bak.get("phone") or "").strip()[:32],
                        op,
                    ),
                )
            conn.commit()
    except UndefinedTable as exc:
        raise HTTPException(status_code=503, detail=f"RL 值班表未就绪：{_DUTY_EXTRAS_SCHEMA_HINT}") from exc
    audit_log("duty.rl_oncall.update", operator=op, count=len(payload.rows))
    return {"ok": True, "count": len(payload.rows)}