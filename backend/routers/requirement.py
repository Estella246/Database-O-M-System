from __future__ import annotations

import json
import urllib.parse
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Any

import psycopg
from psycopg.errors import UndefinedTable

from fastapi import APIRouter, HTTPException, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

from config import (
    _REQUIREMENT_NO_LOCK,
    _REQUIREMENT_SCHEMA_HINT,
    REQUIREMENT_STATUSES,
    REQUIREMENT_CATEGORIES,
    REQUIREMENT_PRIORITIES,
)
from database import db_conn
from models import RequirementCreatePayload, RequirementPatchPayload, RequirementExportPayload
from utils import parse_ymd as _parse_ymd
from whitelist_policy import whitelist_permission_level, whitelist_field_levels

router = APIRouter(prefix="/api/requirements", tags=["requirements"])

# 默认值（与前端/迁移保持一致）
_DEFAULT_CATEGORY = "质量加固和改进"
_DEFAULT_PRIORITY = "中"
_DEFAULT_STATUS = "已接纳"

# 列表/导出/模板列顺序：编号 / 分类 / 代表问题 / 所属领域 / 模块&特性 /
#                        问题描述 / 改进诉求 / 优先级 / 提出人 / 接纳状态 / 计划版本
# (列中文名 -> requirement 表字段名)
_IMPORT_COLUMNS: list[tuple[str, str]] = [
    ("编号", "requirement_no"),
    ("分类", "category"),
    ("代表问题", "represent_issue"),
    ("所属领域", "domain"),
    ("模块&特性", "module_feature"),
    ("问题描述", "description"),
    ("改进诉求", "improvement"),
    ("优先级", "priority"),
    ("提出人", "proposer"),
    ("接纳状态", "status"),
    ("计划版本", "planned_version"),
]


def _display_name_account(conn: psycopg.Connection, account: str) -> str:
    acc = str(account or "").strip()
    if not acc:
        return ""
    row = conn.execute("SELECT user_name FROM user_account WHERE account = %s", (acc,)).fetchone()
    un = str(row["user_name"] or "").strip() if row else ""
    return f"{un} {acc}".strip() if un else acc


def _allocate_requirement_no(conn: psycopg.Connection) -> str:
    """分配自增编号：全局纯数字流水号（max+1）。"""
    conn.execute("SELECT pg_advisory_xact_lock(%s)", (_REQUIREMENT_NO_LOCK,))
    row = conn.execute(
        """
        SELECT COALESCE(MAX(CASE WHEN requirement_no ~ '^[0-9]+$'
                                 THEN CAST(requirement_no AS BIGINT) END), 0) AS mx
        FROM requirement
        """
    ).fetchone()
    return str(int(row["mx"] or 0) + 1)


def _schema_error(exc: UndefinedTable) -> HTTPException:
    return HTTPException(status_code=503, detail=f"质量改进表未就绪：{_REQUIREMENT_SCHEMA_HINT}")


@router.get("")
def list_requirements(
    operator_id: str = "demo_001",
    scope: str = "all",
    status: str = "",
    priority: str = "",
    category: str = "",
    q: str = "",
    page: int = 1,
    page_size: int = 20,
) -> dict:
    op = operator_id.strip() or "demo_001"
    sc = (scope or "all").strip().lower()
    if sc not in ("all", "mine"):
        raise HTTPException(status_code=400, detail="scope 须为 all 或 mine")
    qq = str(q or "").strip()
    status_list = [s.strip() for s in status.split(",") if s.strip()] if status else []
    priority_list = [p.strip() for p in priority.split(",") if p.strip()] if priority else []
    category_list = [c.strip() for c in category.split(",") if c.strip()] if category else []
    pg = max(1, page)
    ps = max(1, min(10000, page_size))
    offset = (pg - 1) * ps
    try:
        with db_conn() as conn:
            where_parts: list[str] = ["1=1"]
            params: list = []
            if sc == "mine":
                where_parts.append("r.creator_id = %s")
                params.append(op)
            if status_list:
                ph = ",".join(["%s"] * len(status_list))
                where_parts.append(f"r.status IN ({ph})")
                params.extend(status_list)
            if priority_list:
                ph = ",".join(["%s"] * len(priority_list))
                where_parts.append(f"r.priority IN ({ph})")
                params.extend(priority_list)
            if category_list:
                ph = ",".join(["%s"] * len(category_list))
                where_parts.append(f"r.category IN ({ph})")
                params.extend(category_list)
            if qq:
                like = f"%{qq}%"
                where_parts.append(
                    "(r.requirement_no ILIKE %s OR r.category ILIKE %s OR r.represent_issue ILIKE %s"
                    " OR r.domain ILIKE %s OR r.module_feature ILIKE %s OR r.description ILIKE %s"
                    " OR r.improvement ILIKE %s OR r.proposer ILIKE %s OR r.status ILIKE %s"
                    " OR r.priority ILIKE %s OR r.planned_version ILIKE %s)"
                )
                params.extend([like] * 11)
            where_sql = " AND ".join(where_parts)
            total = conn.execute(
                f"SELECT COUNT(*) AS cnt FROM requirement r WHERE {where_sql}", tuple(params)
            ).fetchone()["cnt"]
            # 优先级排序：高 < 中 < 低
            rows = conn.execute(
                f"""
                SELECT r.* FROM requirement r
                WHERE {where_sql}
                ORDER BY CASE r.priority WHEN '高' THEN 1 WHEN '中' THEN 2 WHEN '低' THEN 3 ELSE 9 END,
                         r.created_at DESC
                LIMIT %s OFFSET %s
                """,
                tuple(params) + (ps, offset),
            ).fetchall()
    except UndefinedTable as exc:
        raise _schema_error(exc) from exc
    return {"items": [dict(r) for r in rows], "total": int(total or 0), "page": pg, "page_size": ps}


@router.get("/analytics")
def analytics_requirements(
    operator_id: str = "demo_001",
    start_date: str = "",
    end_date: str = "",
    precision: str = "week",
) -> dict:
    _ = operator_id.strip() or "demo_001"
    today = datetime.now().date()
    ed = _parse_ymd(end_date, "end_date") if end_date else today
    sd = _parse_ymd(start_date, "start_date") if start_date else today - timedelta(days=90)
    if sd > ed:
        sd, ed = ed, sd
    prec = (precision or "week").strip().lower()
    if prec not in ("week", "month"):
        raise HTTPException(status_code=400, detail="precision 须为 week 或 month")
    start_dt = datetime(sd.year, sd.month, sd.day, tzinfo=timezone.utc)
    end_dt_exclusive = datetime(ed.year, ed.month, ed.day, tzinfo=timezone.utc) + timedelta(days=1)
    try:
        with db_conn() as conn:
            total = int(conn.execute(
                "SELECT COUNT(*) AS total FROM requirement WHERE created_at >= %s AND created_at < %s",
                (start_dt, end_dt_exclusive),
            ).fetchone()["total"] or 0)

            status_rows = conn.execute(
                "SELECT status, COUNT(*) AS cnt FROM requirement WHERE created_at >= %s AND created_at < %s GROUP BY status",
                (start_dt, end_dt_exclusive),
            ).fetchall()
            by_status = {str(r["status"]): int(r["cnt"]) for r in status_rows}
            status_labels = list(REQUIREMENT_STATUSES)
            status_values = [by_status.get(s, 0) for s in status_labels]
            realized = by_status.get("已实现", 0)
            rejected = by_status.get("拒绝", 0)

            category_rows = conn.execute(
                "SELECT category, COUNT(*) AS cnt FROM requirement WHERE created_at >= %s AND created_at < %s GROUP BY category",
                (start_dt, end_dt_exclusive),
            ).fetchall()
            category_map = {str(r["category"]): int(r["cnt"]) for r in category_rows}
            category_labels = list(REQUIREMENT_CATEGORIES)
            category_values = [category_map.get(c, 0) for c in category_labels]

            prio_rows = conn.execute(
                "SELECT priority, COUNT(*) AS cnt FROM requirement WHERE created_at >= %s AND created_at < %s GROUP BY priority",
                (start_dt, end_dt_exclusive),
            ).fetchall()
            prio_map = {str(r["priority"]): int(r["cnt"]) for r in prio_rows}
            priority_labels = list(REQUIREMENT_PRIORITIES)
            priority_values = [prio_map.get(p, 0) for p in priority_labels]

            trunc = "week" if prec == "week" else "month"
            fmt = "YYYY\"W\"IW" if prec == "week" else "YYYY-MM"
            trend_created_rows = conn.execute(
                "SELECT to_char(date_trunc(%s, created_at), %s) AS label, COUNT(*) AS cnt FROM requirement WHERE created_at >= %s AND created_at < %s GROUP BY label ORDER BY MIN(created_at)",
                (trunc, fmt, start_dt, end_dt_exclusive),
            ).fetchall()
            all_labels = [str(r["label"]) for r in trend_created_rows]
            created_map = {str(r["label"]): int(r["cnt"]) for r in trend_created_rows}

            proposer_rows = conn.execute(
                "SELECT proposer, COUNT(*) AS cnt FROM requirement WHERE created_at >= %s AND created_at < %s GROUP BY proposer ORDER BY cnt DESC LIMIT 10",
                (start_dt, end_dt_exclusive),
            ).fetchall()
    except UndefinedTable as exc:
        raise _schema_error(exc) from exc
    return {
        "kpi": {
            "total": total,
            "by_status": by_status,
            "realized": realized,
            "rejected": rejected,
        },
        "status_distribution": {"labels": status_labels, "values": status_values},
        "category_distribution": {"labels": category_labels, "values": category_values},
        "priority_distribution": {"labels": priority_labels, "values": priority_values},
        "trend": {"labels": all_labels, "created": [created_map.get(l, 0) for l in all_labels]},
        "person_load": {
            "top_proposers": [{"name": str(r["proposer"]), "count": int(r["cnt"])} for r in proposer_rows],
        },
    }


def _validate_enums(category: str, priority: str, status: str) -> None:
    if category not in REQUIREMENT_CATEGORIES:
        raise HTTPException(status_code=400, detail=f"无效分类: {category}")
    if priority not in REQUIREMENT_PRIORITIES:
        raise HTTPException(status_code=400, detail=f"无效优先级: {priority}（须为 高/中/低）")
    if status not in REQUIREMENT_STATUSES:
        raise HTTPException(status_code=400, detail=f"无效接纳状态: {status}")


@router.post("")
def create_requirement(payload: RequirementCreatePayload) -> dict:
    op = payload.operator_id.strip()
    if not op:
        raise HTTPException(status_code=400, detail="operator_id 不能为空")
    if not payload.improvement.strip():
        raise HTTPException(status_code=400, detail="改进诉求不能为空")
    if not payload.proposer.strip():
        raise HTTPException(status_code=400, detail="提出人不能为空")
    cat = payload.category.strip() or _DEFAULT_CATEGORY
    prio = payload.priority.strip() or _DEFAULT_PRIORITY
    st = payload.status.strip() or _DEFAULT_STATUS
    _validate_enums(cat, prio, st)
    try:
        with db_conn() as conn:
            creator_disp = _display_name_account(conn, op)
            req_no = _allocate_requirement_no(conn)
            row = conn.execute(
                """
                INSERT INTO requirement (
                  requirement_no, category, represent_issue, domain, module_feature,
                  description, improvement, priority, proposer, status, planned_version,
                  creator_id, creator_name
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                RETURNING *
                """,
                (
                    req_no, cat, payload.represent_issue.strip(), payload.domain.strip(),
                    payload.module_feature.strip(), payload.description.strip(),
                    payload.improvement.strip(), prio, payload.proposer.strip(), st,
                    payload.planned_version.strip(), op, creator_disp,
                ),
            ).fetchone()
            if not row:
                raise HTTPException(status_code=500, detail="写入质量改进失败")
            req_id = int(row["id"])
            conn.execute(
                "INSERT INTO requirement_log (requirement_id, action, to_status, operator_id, operator_name, comment) VALUES (%s, %s, %s, %s, %s, %s)",
                (req_id, "created", st, op, creator_disp, ""),
            )
            conn.commit()
    except HTTPException:
        raise
    except UndefinedTable as exc:
        raise _schema_error(exc) from exc
    return dict(row)


@router.get("/{req_id:int}")
def get_requirement(req_id: int, operator_id: str = "demo_001") -> dict:
    _ = operator_id
    try:
        with db_conn() as conn:
            row = conn.execute("SELECT * FROM requirement WHERE id = %s", (req_id,)).fetchone()
    except UndefinedTable as exc:
        raise _schema_error(exc) from exc
    if not row:
        raise HTTPException(status_code=404, detail="需求不存在")
    return dict(row)


@router.patch("/{req_id:int}")
def patch_requirement(req_id: int, payload: RequirementPatchPayload) -> dict:
    op = payload.operator_id.strip() or "demo_001"
    try:
        with db_conn() as conn:
            row = conn.execute("SELECT * FROM requirement WHERE id = %s FOR UPDATE", (req_id,)).fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="需求不存在")
            old = dict(row)
            updates: dict[str, Any] = {}
            changed: dict[str, list] = {}
            # 普通文本字段
            text_fields = {
                "represent_issue": payload.represent_issue,
                "domain": payload.domain,
                "module_feature": payload.module_feature,
                "description": payload.description,
                "improvement": payload.improvement,
                "proposer": payload.proposer,
                "planned_version": payload.planned_version,
            }
            for field, val in text_fields.items():
                if val is not None:
                    new_val = str(val).strip()
                    old_val = str(old.get(field, "") or "").strip()
                    if new_val != old_val:
                        updates[field] = new_val
                        changed[field] = [old_val, new_val]
            # 枚举字段
            enum_fields = [
                ("category", payload.category, REQUIREMENT_CATEGORIES, "无效分类"),
                ("priority", payload.priority, REQUIREMENT_PRIORITIES, "无效优先级"),
                ("status", payload.status, REQUIREMENT_STATUSES, "无效接纳状态"),
            ]
            from_status = None
            to_status = None
            for field, val, allowed, errlabel in enum_fields:
                if val is None:
                    continue
                new_val = str(val).strip()
                if new_val not in allowed:
                    raise HTTPException(status_code=400, detail=f"{errlabel}: {new_val}")
                old_val = str(old.get(field, "") or "").strip()
                if new_val != old_val:
                    updates[field] = new_val
                    changed[field] = [old_val, new_val]
                    if field == "status":
                        from_status, to_status = old_val, new_val
            if not updates:
                conn.rollback()
                return dict(old)
            set_parts = [f"{k} = %s" for k in updates]
            set_parts.append("updated_at = NOW()")
            vals = list(updates.values())
            vals.append(req_id)
            conn.execute(f"UPDATE requirement SET {', '.join(set_parts)} WHERE id = %s", tuple(vals))
            operator_disp = _display_name_account(conn, op)
            action = "status_changed" if to_status else "updated"
            conn.execute(
                """
                INSERT INTO requirement_log (requirement_id, action, from_status, to_status, changed_fields, operator_id, operator_name, comment)
                VALUES (%s, %s, %s, %s, %s::jsonb, %s, %s, %s)
                """,
                (req_id, action, from_status, to_status,
                 psycopg.types.json.Jsonb(changed) if changed else None,
                 op, operator_disp, payload.comment.strip()),
            )
            conn.commit()
            new_row = conn.execute("SELECT * FROM requirement WHERE id = %s", (req_id,)).fetchone()
            return dict(new_row)
    except HTTPException:
        raise
    except UndefinedTable as exc:
        raise _schema_error(exc) from exc


@router.get("/{req_id:int}/logs")
def get_requirement_logs(req_id: int, operator_id: str = "demo_001") -> dict:
    _ = operator_id
    try:
        with db_conn() as conn:
            rows = conn.execute(
                """
                SELECT id, action, from_status, to_status, changed_fields, comment,
                       operator_id, operator_name, created_at
                FROM requirement_log
                WHERE requirement_id = %s
                ORDER BY created_at DESC
                """,
                (req_id,),
            ).fetchall()
    except UndefinedTable as exc:
        raise _schema_error(exc) from exc
    return {"items": [dict(r) for r in rows]}


@router.delete("/{req_id:int}")
def delete_requirement(req_id: int, operator_id: str = "demo_001") -> dict:
    op = operator_id.strip() or "demo_001"
    try:
        with db_conn() as conn:
            row = conn.execute("SELECT creator_id FROM requirement WHERE id = %s", (req_id,)).fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="需求不存在")
            if str(row["creator_id"] or "").strip() != op:
                raise HTTPException(status_code=403, detail="仅创建人可删除")
            conn.execute("DELETE FROM requirement WHERE id = %s", (req_id,))
            conn.commit()
    except HTTPException:
        raise
    except UndefinedTable as exc:
        raise _schema_error(exc) from exc
    return {"deleted": req_id}


def _excel_header_style(ws, headers: list[str]) -> Border:
    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"),
                         top=Side(style="thin"), bottom=Side(style="thin"))
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
    return thin_border


@router.post("/export")
def export_requirements(payload: RequirementExportPayload) -> StreamingResponse:
    """导出质量改进为 Excel 文件。"""
    op = payload.operator_id.strip() or "demo_001"
    cols_sql = ", ".join(field for _, field in _IMPORT_COLUMNS)
    try:
        with db_conn() as conn:
            wl = whitelist_field_levels(conn, op)
            if whitelist_permission_level(wl, "requirement_export") == "hidden":
                raise HTTPException(status_code=403, detail="无导出权限")
            rows = conn.execute(
                f"""
                SELECT {cols_sql}, creator_name, created_at, updated_at
                FROM requirement
                ORDER BY CASE priority WHEN '高' THEN 1 WHEN '中' THEN 2 WHEN '低' THEN 3 ELSE 9 END,
                         created_at DESC
                """
            ).fetchall()
    except UndefinedTable as exc:
        raise _schema_error(exc) from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "质量改进导出"
    headers = [name for name, _ in _IMPORT_COLUMNS] + ["创建人", "创建时间", "更新时间"]
    thin_border = _excel_header_style(ws, headers)

    for row_idx, row in enumerate(rows, start=2):
        created_at_val = row.get("created_at")
        updated_at_val = row.get("updated_at")
        values = [str(row.get(field) or "") for _, field in _IMPORT_COLUMNS]
        values.append(str(row.get("creator_name") or ""))
        values.append(created_at_val.strftime("%Y-%m-%d %H:%M:%S") if created_at_val else "")
        values.append(updated_at_val.strftime("%Y-%m-%d %H:%M:%S") if updated_at_val else "")
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    today = datetime.now().strftime("%Y-%m-%d")
    filename = f"quality_improvement_{op}_{today}.xlsx"
    encoded_filename = urllib.parse.quote(f"质量改进导出_{op}_{today}.xlsx", safe="")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{filename}\"; filename*=UTF-8''{encoded_filename}"},
    )


@router.get("/import-template")
def get_import_template(operator_id: str = "demo_001") -> StreamingResponse:
    """下载质量改进导入模板 Excel 文件。"""
    op = operator_id.strip() or "demo_001"
    try:
        with db_conn() as conn:
            wl = whitelist_field_levels(conn, op)
            if whitelist_permission_level(wl, "requirement_import") == "hidden":
                raise HTTPException(status_code=403, detail="无导入权限")
    except UndefinedTable as exc:
        raise _schema_error(exc) from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "质量改进导入模板"
    headers = [name for name, _ in _IMPORT_COLUMNS]
    thin_border = _excel_header_style(ws, headers)
    # 示例行（编号留空=新增）
    example_values = [
        "", "质量加固和改进", "YW20260525001 主备倒换异常", "存储引擎", "空间管理/回收站",
        "回收站空间未及时回收导致磁盘满", "增加后台自动回收与水位告警", "高", "张三 zhangsan",
        "已接纳", "V8.2.0",
    ]
    for col_idx, value in enumerate(example_values, start=1):
        cell = ws.cell(row=2, column=col_idx, value=value)
        cell.border = thin_border

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    encoded_filename = urllib.parse.quote("质量改进导入模板.xlsx", safe="")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"quality_improvement_template.xlsx\"; filename*=UTF-8''{encoded_filename}"},
    )


def _parse_excel_import(file_content: bytes) -> tuple[list[dict], list[dict]]:
    """解析 Excel 导入文件，返回 (数据行列表, 错误列表)。"""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(file_content))
    ws = wb.active
    errors: list[dict] = []
    headers: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        header_val = ws.cell(row=1, column=col).value
        if header_val:
            headers[str(header_val).strip()] = col
    expected = [name for name, _ in _IMPORT_COLUMNS]
    for h in expected:
        if h not in headers:
            errors.append({"row": 1, "field": "表头", "message": f"缺少必填列：{h}"})
    if errors:
        return [], errors
    rows: list[dict] = []
    # 第 3 行起（跳过表头与示例行）
    for row_idx in range(3, ws.max_row + 1):
        row_data: dict[str, Any] = {}
        for h, col in headers.items():
            val = ws.cell(row=row_idx, column=col).value
            row_data[h] = val if val is not None else ""
        # 改进诉求为空视为空行跳过
        if not str(row_data.get("改进诉求", "")).strip():
            continue
        row_data["_row_idx"] = row_idx
        rows.append(row_data)
    return rows, errors


@router.post("/import")
async def import_requirements(
    file: UploadFile = File(...),
    operator_id: str = Form(...),
) -> dict:
    """批量导入质量改进。编号为空=新增；填写已有编号=更新。"""
    op = operator_id.strip() or "demo_001"
    try:
        with db_conn() as conn:
            wl = whitelist_field_levels(conn, op)
            if whitelist_permission_level(wl, "requirement_import") == "hidden":
                raise HTTPException(status_code=403, detail="无导入权限")
    except UndefinedTable as exc:
        raise _schema_error(exc) from exc

    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 格式文件")

    try:
        content = await file.read()
        rows, parse_errors = _parse_excel_import(content)
        if parse_errors:
            raise HTTPException(status_code=400, detail=json.dumps({"success": False, "error_type": "validation_failed", "errors": parse_errors}))
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail="文件无法解析，请检查文件格式") from e

    if not rows:
        return {"success": True, "total": 0, "created": 0, "updated": 0, "message": "导入成功，共0条"}

    validation_errors: list[dict] = []
    created_count = 0
    updated_count = 0
    try:
        with db_conn() as conn:
            existing: dict[str, int] = {}
            for r in conn.execute("SELECT id, requirement_no FROM requirement").fetchall():
                no = str(r["requirement_no"] or "").strip()
                if no:
                    existing[no] = int(r["id"])
            operator_disp = _display_name_account(conn, op)

            # 全量宽松校验：枚举无效回落默认值；必填(改进诉求/提出人)缺失报错
            for row_data in rows:
                row_idx = row_data["_row_idx"]
                improvement = str(row_data.get("改进诉求", "")).strip()
                proposer = str(row_data.get("提出人", "")).strip()
                if not improvement:
                    validation_errors.append({"row": row_idx, "field": "改进诉求", "message": "必填字段不能为空"})
                if not proposer:
                    validation_errors.append({"row": row_idx, "field": "提出人", "message": "必填字段不能为空"})
                category = str(row_data.get("分类", "")).strip() or _DEFAULT_CATEGORY
                if category not in REQUIREMENT_CATEGORIES:
                    category = _DEFAULT_CATEGORY
                priority = str(row_data.get("优先级", "")).strip() or _DEFAULT_PRIORITY
                if priority not in REQUIREMENT_PRIORITIES:
                    priority = _DEFAULT_PRIORITY
                status = str(row_data.get("接纳状态", "")).strip() or _DEFAULT_STATUS
                if status not in REQUIREMENT_STATUSES:
                    status = _DEFAULT_STATUS
                req_no = str(row_data.get("编号", "")).strip()
                row_data["_validated"] = {
                    "category": category,
                    "represent_issue": str(row_data.get("代表问题", "")).strip(),
                    "domain": str(row_data.get("所属领域", "")).strip(),
                    "module_feature": str(row_data.get("模块&特性", "")).strip(),
                    "description": str(row_data.get("问题描述", "")).strip(),
                    "improvement": improvement,
                    "priority": priority,
                    "proposer": proposer,
                    "status": status,
                    "planned_version": str(row_data.get("计划版本", "")).strip(),
                    "req_no": req_no,
                    "existing_id": existing.get(req_no) if req_no else None,
                }

            if validation_errors:
                raise HTTPException(status_code=400, detail=json.dumps({"success": False, "error_type": "validation_failed", "errors": validation_errors}))

            for row_data in rows:
                v = row_data["_validated"]
                if v["existing_id"]:
                    req_id = v["existing_id"]
                    old_row = conn.execute("SELECT * FROM requirement WHERE id = %s FOR UPDATE", (req_id,)).fetchone()
                    old = dict(old_row)
                    updates: dict[str, Any] = {}
                    changed: dict[str, list] = {}
                    fields = ["category", "represent_issue", "domain", "module_feature",
                              "description", "improvement", "priority", "proposer",
                              "status", "planned_version"]
                    for f in fields:
                        new_val = v[f]
                        old_val = str(old.get(f, "") or "").strip()
                        if new_val != old_val:
                            updates[f] = new_val
                            changed[f] = [old_val, new_val]
                    if updates:
                        set_parts = [f"{k} = %s" for k in updates]
                        set_parts.append("updated_at = NOW()")
                        vals = list(updates.values()) + [req_id]
                        conn.execute(f"UPDATE requirement SET {', '.join(set_parts)} WHERE id = %s", tuple(vals))
                        to_status = updates.get("status")
                        from_status = changed["status"][0] if "status" in changed else None
                        action = "status_changed" if to_status else "updated"
                        conn.execute(
                            "INSERT INTO requirement_log (requirement_id, action, from_status, to_status, changed_fields, operator_id, operator_name, comment) VALUES (%s, %s, %s, %s, %s::jsonb, %s, %s, %s)",
                            (req_id, action, from_status, to_status, psycopg.types.json.Jsonb(changed), op, operator_disp, ""),
                        )
                        updated_count += 1
                else:
                    req_no_final = _allocate_requirement_no(conn)
                    new_row = conn.execute(
                        """
                        INSERT INTO requirement (
                          requirement_no, category, represent_issue, domain, module_feature,
                          description, improvement, priority, proposer, status, planned_version,
                          creator_id, creator_name
                        )
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                        RETURNING id
                        """,
                        (req_no_final, v["category"], v["represent_issue"], v["domain"],
                         v["module_feature"], v["description"], v["improvement"], v["priority"],
                         v["proposer"], v["status"], v["planned_version"], op, operator_disp),
                    ).fetchone()
                    conn.execute(
                        "INSERT INTO requirement_log (requirement_id, action, to_status, operator_id, operator_name, comment) VALUES (%s, %s, %s, %s, %s, %s)",
                        (new_row["id"], "created", v["status"], op, operator_disp, ""),
                    )
                    created_count += 1
            conn.commit()
    except HTTPException:
        raise
    except UndefinedTable as exc:
        raise _schema_error(exc) from exc
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"导入失败：{str(e)}") from e

    total = created_count + updated_count
    return {
        "success": True,
        "total": total,
        "created": created_count,
        "updated": updated_count,
        "message": f"成功导入{total}条（新增{created_count}条，更新{updated_count}条）",
    }
