"""M10 质量改进（原需求管理）接口测试（迁移 0083 后的新模型）。

新模型字段：编号(自增) / 分类 / 代表问题 / 所属领域 / 模块&特性 / 问题描述 /
            改进诉求 / 优先级(高/中/低) / 提出人 / 接纳状态 / 计划版本。
覆盖：增删改查、搜索/筛选/分页、日志、分析看板、导入/导出/模板。
"""
import io

import pytest

OP = "admin"


@pytest.fixture(scope="module", autouse=True)
def _ensure_req_whitelist(api_client):
    """确保 admin 角色拥有质量改进导入/导出/创建白名单（export 在策略表中是受管键）。"""
    api_client.post("/api/admin/permissions/bulk", json={
        "operator_id": "admin",
        "items": [
            {"role_code": "admin", "is_pl": False, "node_key": "__whitelist__",
             "field_key": f, "permission_level": "readonly"}
            for f in ("requirement_export", "requirement_import", "requirement_create")
        ],
    })
    yield


def _create(api_client, operator_id=OP, **overrides):
    payload = {
        "operator_id": operator_id,
        "category": "质量加固和改进",
        "represent_issue": "YW20260525001 主备倒换异常",
        "domain": "存储引擎",
        "module_feature": "空间管理/回收站",
        "description": "回收站空间未及时回收导致磁盘满",
        "improvement": "增加后台自动回收与水位告警",
        "priority": "高",
        "proposer": "张三 zhangsan",
        "status": "已接纳",
        "planned_version": "V8.2.0",
    }
    payload.update(overrides)
    return api_client.post("/api/requirements", json=payload)


class TestRequirementCreate:
    def test_tc_m10_001_create(self, api_client):
        r = _create(api_client, improvement="改进-001")
        assert r.status_code == 200, r.text
        b = r.json()
        assert b["improvement"] == "改进-001"
        assert b["category"] == "质量加固和改进"
        assert b["priority"] == "高"
        assert b["status"] == "已接纳"
        assert b["requirement_no"]  # 自增编号非空

    def test_tc_m10_002_auto_increment_no(self, api_client):
        a = _create(api_client, improvement="自增A").json()
        b = _create(api_client, improvement="自增B").json()
        assert int(b["requirement_no"]) == int(a["requirement_no"]) + 1

    def test_tc_m10_003_missing_improvement(self, api_client):
        r = _create(api_client, improvement="   ")
        assert r.status_code == 400

    def test_tc_m10_004_missing_proposer(self, api_client):
        r = _create(api_client, proposer="")
        assert r.status_code == 400

    def test_tc_m10_005_invalid_category(self, api_client):
        r = _create(api_client, category="管控需求")
        assert r.status_code == 400

    def test_tc_m10_006_invalid_priority(self, api_client):
        r = _create(api_client, priority="P1")
        assert r.status_code == 400

    def test_tc_m10_007_invalid_status(self, api_client):
        r = _create(api_client, status="开发中")
        assert r.status_code == 400

    def test_tc_m10_008_defaults_applied(self, api_client):
        r = api_client.post("/api/requirements", json={
            "operator_id": OP, "improvement": "仅必填", "proposer": "李四 lisi",
        })
        assert r.status_code == 200, r.text
        b = r.json()
        assert b["category"] == "质量加固和改进"
        assert b["priority"] == "中"
        assert b["status"] == "待评审"
        assert b["proposed_at"]  # 未填时默认当天

    def test_tc_m10_009_proposed_at_explicit(self, api_client):
        r = _create(api_client, improvement="带提出时间", proposed_at="2026-03-15")
        assert r.status_code == 200, r.text
        assert str(r.json()["proposed_at"])[:10] == "2026-03-15"


class TestRequirementList:
    def test_tc_m10_010_list_all(self, api_client):
        _create(api_client, improvement="列表项")
        r = api_client.get("/api/requirements", params={"operator_id": OP})
        assert r.status_code == 200
        assert r.json()["total"] >= 1

    def test_tc_m10_011_list_mine(self, api_client):
        _create(api_client, operator_id=OP, improvement="我的项")
        r = api_client.get("/api/requirements", params={"operator_id": OP, "scope": "mine"})
        assert r.status_code == 200

    def test_tc_m10_012_invalid_scope(self, api_client):
        r = api_client.get("/api/requirements", params={"operator_id": OP, "scope": "assigned"})
        assert r.status_code == 400

    def test_tc_m10_013_search(self, api_client):
        _create(api_client, improvement="搜索关键词_zzx", domain="网络子系统")
        r = api_client.get("/api/requirements", params={"operator_id": OP, "q": "zzx"})
        assert r.status_code == 200
        assert any("zzx" in str(it.get("improvement", "")) for it in r.json()["items"])

    def test_tc_m10_014_filter_priority(self, api_client):
        _create(api_client, improvement="低优先", priority="低")
        r = api_client.get("/api/requirements", params={"operator_id": OP, "priority": "低", "page_size": 100})
        assert r.status_code == 200
        assert all(it["priority"] == "低" for it in r.json()["items"])

    def test_tc_m10_015_filter_category_and_status(self, api_client):
        _create(api_client, improvement="分类筛选", category="测试加固", status="拒绝")
        r = api_client.get("/api/requirements", params={"operator_id": OP, "category": "测试加固", "status": "拒绝", "page_size": 100})
        assert r.status_code == 200
        for it in r.json()["items"]:
            assert it["category"] == "测试加固" and it["status"] == "拒绝"

    def test_tc_m10_016_pagination(self, api_client):
        for i in range(3):
            _create(api_client, improvement=f"分页{i}")
        r = api_client.get("/api/requirements", params={"operator_id": OP, "page": 1, "page_size": 2})
        assert r.status_code == 200
        assert len(r.json()["items"]) <= 2


class TestRequirementDetailPatchDelete:
    def test_tc_m10_020_get_detail(self, api_client):
        rid = _create(api_client, improvement="详情项").json()["id"]
        r = api_client.get(f"/api/requirements/{rid}", params={"operator_id": OP})
        assert r.status_code == 200
        assert r.json()["id"] == rid

    def test_tc_m10_021_get_nonexistent(self, api_client):
        r = api_client.get("/api/requirements/99999999", params={"operator_id": OP})
        assert r.status_code == 404

    def test_tc_m10_022_patch_fields(self, api_client):
        rid = _create(api_client, improvement="待改").json()["id"]
        r = api_client.patch(f"/api/requirements/{rid}", json={
            "operator_id": OP, "improvement": "已改", "domain": "优化器", "priority": "中",
        })
        assert r.status_code == 200
        b = r.json()
        assert b["improvement"] == "已改" and b["domain"] == "优化器" and b["priority"] == "中"

    def test_tc_m10_023_patch_status_free(self, api_client):
        # 接纳状态可直接改为任意合法值（无流转限制）
        rid = _create(api_client, status="已接纳").json()["id"]
        r = api_client.patch(f"/api/requirements/{rid}", json={"operator_id": OP, "status": "已实现"})
        assert r.status_code == 200
        assert r.json()["status"] == "已实现"

    def test_tc_m10_024_patch_invalid_enum(self, api_client):
        rid = _create(api_client).json()["id"]
        r = api_client.patch(f"/api/requirements/{rid}", json={"operator_id": OP, "priority": "紧急"})
        assert r.status_code == 400

    def test_tc_m10_025_patch_nonexistent(self, api_client):
        r = api_client.patch("/api/requirements/99999999", json={"operator_id": OP, "improvement": "x"})
        assert r.status_code == 404

    def test_tc_m10_026_logs(self, api_client):
        rid = _create(api_client).json()["id"]
        api_client.patch(f"/api/requirements/{rid}", json={"operator_id": OP, "status": "部分接纳"})
        r = api_client.get(f"/api/requirements/{rid}/logs", params={"operator_id": OP})
        assert r.status_code == 200
        actions = [lg["action"] for lg in r.json()["items"]]
        assert "created" in actions and "status_changed" in actions

    def test_tc_m10_027_delete_by_creator(self, api_client):
        rid = _create(api_client, operator_id=OP).json()["id"]
        r = api_client.delete(f"/api/requirements/{rid}", params={"operator_id": OP})
        assert r.status_code == 200

    def test_tc_m10_028_delete_non_creator(self, api_client):
        rid = _create(api_client, operator_id=OP).json()["id"]
        r = api_client.delete(f"/api/requirements/{rid}", params={"operator_id": "someone_else"})
        assert r.status_code == 403

    def test_tc_m10_029_delete_nonexistent(self, api_client):
        r = api_client.delete("/api/requirements/99999999", params={"operator_id": OP})
        assert r.status_code == 404


class TestRequirementAnalytics:
    def test_tc_m10_030_analytics_shape(self, api_client):
        _create(api_client, improvement="分析项", priority="高", category="测试加固", status="已实现")
        r = api_client.get("/api/requirements/analytics", params={"operator_id": OP})
        assert r.status_code == 200
        d = r.json()
        assert d["status_distribution"]["labels"] == ["待评审", "已实现", "已接纳", "部分接纳", "拒绝"]
        assert d["priority_distribution"]["labels"] == ["高", "中", "低"]
        assert d["category_distribution"]["labels"] == ["定位定界", "测试加固", "快速恢复", "需求", "质量加固和改进"]
        assert "top_proposers" in d["person_load"]
        assert "created" in d["trend"]

    def test_tc_m10_031_analytics_invalid_precision(self, api_client):
        r = api_client.get("/api/requirements/analytics", params={"operator_id": OP, "precision": "day"})
        assert r.status_code == 400

    def test_tc_m10_032_analytics_precision_month(self, api_client):
        r = api_client.get("/api/requirements/analytics", params={"operator_id": OP, "precision": "month"})
        assert r.status_code == 200


_IMPORT_HEADERS = ["编号", "分类", "代表问题", "所属领域", "模块&特性", "问题描述",
                   "改进诉求", "优先级", "提出人", "提出时间", "接纳状态", "计划版本"]


def _build_xlsx(data_rows, example_row=True):
    """构造导入用 xlsx：第1行表头，第2行起为数据。

    example_row=True 时在第 2 行插入空示例行（「改进诉求」为空，导入时被跳过），
    模拟模板布局；example_row=False 时数据直接从第 2 行开始，模拟导出文件。
    """
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(_IMPORT_HEADERS)
    if example_row:
        ws.append([""] * len(_IMPORT_HEADERS))  # 空示例行（改进诉求为空→跳过）
    for row in data_rows:
        ws.append([row.get(h, "") for h in _IMPORT_HEADERS])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class TestRequirementImportExport:
    def test_tc_m10_040_template_download(self, api_client):
        r = api_client.get("/api/requirements/import-template", params={"operator_id": OP})
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")

    def test_tc_m10_041_export(self, api_client):
        _create(api_client, improvement="导出项")
        r = api_client.post("/api/requirements/export", json={"operator_id": OP})
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")

    def test_tc_m10_042_import_create(self, api_client):
        buf = _build_xlsx([
            {"分类": "测试加固", "代表问题": "imp-rep", "所属领域": "SQL引擎",
             "模块&特性": "优化器", "问题描述": "慢查询", "改进诉求": "导入新增项A",
             "优先级": "中", "提出人": "王五 wangwu", "接纳状态": "已接纳", "计划版本": "V9"},
        ])
        r = api_client.post(
            "/api/requirements/import",
            files={"file": ("t.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"operator_id": OP},
        )
        assert r.status_code == 200, r.text
        assert r.json()["created"] == 1

    def test_tc_m10_043_import_lenient_enum(self, api_client):
        # 非法枚举回落默认值，不报错
        buf = _build_xlsx([
            {"分类": "乱填", "改进诉求": "导入宽松项", "优先级": "超高",
             "提出人": "赵六 zhaoliu", "接纳状态": "未知"},
        ])
        r = api_client.post(
            "/api/requirements/import",
            files={"file": ("t.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"operator_id": OP},
        )
        assert r.status_code == 200, r.text
        assert r.json()["created"] == 1
        lst = api_client.get("/api/requirements", params={"operator_id": OP, "q": "导入宽松项", "page_size": 5}).json()
        it = lst["items"][0]
        assert it["category"] == "质量加固和改进" and it["priority"] == "中" and it["status"] == "待评审"

    def test_tc_m10_044_import_update_by_no(self, api_client):
        created = _create(api_client, improvement="原始诉求").json()
        no = created["requirement_no"]
        buf = _build_xlsx([
            {"编号": no, "分类": "快速恢复", "改进诉求": "更新后的诉求",
             "优先级": "低", "提出人": "张三 zhangsan", "接纳状态": "部分接纳"},
        ])
        r = api_client.post(
            "/api/requirements/import",
            files={"file": ("t.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"operator_id": OP},
        )
        assert r.status_code == 200, r.text
        assert r.json()["updated"] == 1
        detail = api_client.get(f"/api/requirements/{created['id']}", params={"operator_id": OP}).json()
        assert detail["improvement"] == "更新后的诉求" and detail["priority"] == "低" and detail["status"] == "部分接纳"

    def test_tc_m10_045_import_missing_required(self, api_client):
        # 改进诉求非空但提出人为空 → 校验报错
        buf = _build_xlsx([
            {"改进诉求": "有诉求无提出人", "提出人": ""},
        ])
        r = api_client.post(
            "/api/requirements/import",
            files={"file": ("t.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"operator_id": OP},
        )
        assert r.status_code == 400

    def test_tc_m10_046_import_non_xlsx(self, api_client):
        r = api_client.post(
            "/api/requirements/import",
            files={"file": ("t.txt", io.BytesIO(b"hi"), "text/plain")},
            data={"operator_id": OP},
        )
        assert r.status_code == 400

    def test_tc_m10_047_import_data_from_row2(self, api_client):
        # 导出文件无示例行：数据从第 2 行开始，应全部导入（不丢首行）
        buf = _build_xlsx([
            {"分类": "测试加固", "改进诉求": "导出回导首行", "优先级": "高",
             "提出人": "钱七 qianqi", "接纳状态": "已接纳"},
            {"分类": "测试加固", "改进诉求": "导出回导次行", "优先级": "中",
             "提出人": "钱七 qianqi", "接纳状态": "已接纳"},
        ], example_row=False)
        r = api_client.post(
            "/api/requirements/import",
            files={"file": ("t.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"operator_id": OP},
        )
        assert r.status_code == 200, r.text
        assert r.json()["created"] == 2

    def test_tc_m10_048_export_roundtrip_importable(self, api_client):
        # 导出文件列与模板一致（无创建人/时间戳列），可直接再导入
        _create(api_client, improvement="导出回导项")
        exp = api_client.post("/api/requirements/export", json={"operator_id": OP})
        assert exp.status_code == 200
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(exp.content)).active
        header = [c.value for c in ws[1]]
        assert header == _IMPORT_HEADERS  # 不含 创建人/创建时间/更新时间
        r = api_client.post(
            "/api/requirements/import",
            files={"file": ("exp.xlsx", io.BytesIO(exp.content), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"operator_id": OP},
        )
        assert r.status_code == 200, r.text
        # 导出含已有编号 → 全部按更新处理，不应丢行也不应报错
        assert r.json()["updated"] >= 1 and r.json()["created"] == 0


class TestRequirementWhitelist:
    def _set_hidden(self, api_client, field_key: str):
        api_client.post("/api/admin/permissions/bulk", json={
            "operator_id": "admin",
            "items": [{
                "role_code": "admin",
                "is_pl": False,
                "node_key": "__whitelist__",
                "field_key": field_key,
                "permission_level": "hidden",
            }],
        })

    def _restore_readonly(self, api_client, field_keys):
        api_client.post("/api/admin/permissions/bulk", json={
            "operator_id": "admin",
            "items": [
                {
                    "role_code": "admin",
                    "is_pl": False,
                    "node_key": "__whitelist__",
                    "field_key": f,
                    "permission_level": "readonly",
                }
                for f in field_keys
            ],
        })

    def test_tc_m10_049_list_denied_when_requirement_list_hidden(self, api_client):
        self._set_hidden(api_client, "requirement_list")
        try:
            r = api_client.get("/api/requirements", params={"operator_id": OP})
            assert r.status_code == 403
        finally:
            self._restore_readonly(api_client, ["requirement_list"])

    def test_tc_m10_050_create_denied_when_requirement_create_hidden(self, api_client):
        self._set_hidden(api_client, "requirement_create")
        try:
            r = _create(api_client, improvement="无新建权限")
            assert r.status_code == 403
        finally:
            self._restore_readonly(api_client, ["requirement_create"])

    def test_tc_m10_051_export_denied_when_requirement_export_hidden(self, api_client):
        self._set_hidden(api_client, "requirement_export")
        try:
            r = api_client.post("/api/requirements/export", json={"operator_id": OP})
            assert r.status_code == 403
        finally:
            self._restore_readonly(api_client, ["requirement_export"])

    def test_tc_m10_052_import_template_denied_when_requirement_import_hidden(self, api_client):
        self._set_hidden(api_client, "requirement_import")
        try:
            r = api_client.get("/api/requirements/import-template", params={"operator_id": OP})
            assert r.status_code == 403
        finally:
            self._restore_readonly(api_client, ["requirement_import"])
